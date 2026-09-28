"""在隔离的本机演示服务上走真实资料/助手/审核链路，保留可复核记录。

只允许 127.0.0.1:18765，读取 data/alpha-demo 本地虚构账号；不使用开发业务数据库。
"""

import io
import json
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from uuid import uuid4

import httpx
from docx import Document


def main():
    output = Path("artifacts/alpha_handoff")
    output.mkdir(parents=True, exist_ok=True)
    credentials = json.loads(
        Path("data/alpha-demo/local-credentials.json").read_text(encoding="utf-8")
    )
    evidence = {"status": "running", "steps": []}
    with httpx.Client(base_url="http://127.0.0.1:18765", timeout=150, trust_env=False) as client:

        def call(method, path, **kwargs):
            response = client.request(method, path, **kwargs)
            if response.status_code >= 400:
                raise RuntimeError(
                    f"{method} {path}: {response.status_code} {response.json().get('error', {}).get('code')}"
                )
            return response.json()

        token = call(
            "POST",
            "/api/v1/auth/login",
            json={"login_name": "demo1", "password": credentials["password"]},
        )["access_token"]
        client.headers["Authorization"] = "Bearer " + token
        pid = call("GET", "/api/v1/projects")["items"][0]["id"]
        prefix = f"/api/v1/projects/{pid}"

        def wait_document(did):
            for _ in range(120):
                doc = call("GET", prefix + f"/documents/{did}")
                if doc["parse_status"] == doc["index_status"] == "ready":
                    return doc
                if "failed" in {doc["parse_status"], doc["index_status"]}:
                    raise RuntimeError("document job failed: " + str(doc.get("error_message")))
                time.sleep(1)
            raise RuntimeError("document job did not finish within 120s")

        def ask(entry, text):
            key = uuid4().hex
            with ThreadPoolExecutor() as pool:
                future = pool.submit(
                    call,
                    "POST",
                    prefix + "/assistant/messages",
                    json={"entry": entry, "text": text, "request_key": key},
                )
                seen = set()
                while not future.done():
                    response = client.get(
                        prefix + "/assistant/progress", params={"request_key": key}
                    )
                    if response.status_code == 200:
                        state = response.json()
                        with (output / "progress_events.jsonl").open("a", encoding="utf-8") as file:
                            for event in state["events"]:
                                identity = (event["seq"], event["state"])
                                if identity not in seen:
                                    file.write(
                                        json.dumps(
                                            {"run_id": state["run_id"], **event}, ensure_ascii=False
                                        )
                                        + "\n"
                                    )
                                    seen.add(identity)
                    time.sleep(0.3)
                result = future.result()
            final = call("GET", prefix + "/assistant/progress", params={"request_key": key})
            with (output / "progress_events.jsonl").open("a", encoding="utf-8") as file:
                for event in final["events"]:
                    if (event["seq"], event["state"]) not in seen:
                        file.write(
                            json.dumps({"run_id": final["run_id"], **event}, ensure_ascii=False)
                            + "\n"
                        )
            evidence["steps"].append({"kind": entry, "result": result, "progress": final})
            return result

        try:
            task = call(
                "POST",
                prefix + "/tasks",
                json={
                    "title": "Alpha登录接口",
                    "aliases": ["Alpha认证"],
                    "progress": 20,
                    "status": "in_progress",
                },
            )
            docx = Document()
            docx.add_paragraph("2026-09-29 会议记录：以下是待审核的进度更新。")
            table = docx.add_table(rows=2, cols=3)
            for cell, value in zip(table.rows[0].cells, ["任务", "进度", "备注"]):
                cell.text = value
            for cell, value in zip(
                table.rows[1].cells,
                [f"Alpha登录接口 TASK-{task['id']}", "70%", "负责人和截止日不变"],
            ):
                cell.text = value
            docx.add_paragraph("新资料认为支付回调失败最多重试3次，尚待正式确认。")
            stream = io.BytesIO()
            docx.save(stream)
            uploaded = call(
                "POST",
                prefix + "/documents",
                files={
                    "file": (
                        "Alpha会议表格.docx",
                        stream.getvalue(),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
            )
            did = uploaded["document"]["id"]
            doc = wait_document(did)
            blocks = call("GET", prefix + f"/documents/{did}/blocks", params={"version": 1})[
                "items"
            ]
            assert any("70%" in b["text"] and "table:" in (b.get("locator") or "") for b in blocks)
            evidence["steps"].append({"kind": "table_ingest", "document": doc, "blocks": blocks})
            prior = call(
                "POST",
                prefix + "/documents",
                files={
                    "file": (
                        "Alpha旧版约定.txt",
                        "旧资料认为支付回调失败最多重试5次，该方案未确认。".encode(),
                        "text/plain",
                    )
                },
            )
            wait_document(prior["document"]["id"])
            answer = ask(
                "project_assistant",
                "两份资料对支付回调失败后的重试次数分别怎么说？尚未确认时不要替我选择一个。",
            )
            assert len({ref["document_id"] for ref in answer["evidence"]}) >= 2
            suggestion = ask(
                "task_assistant",
                f"把 TASK-{task['id']} 的进度改为70%，状态为进行中，其他字段不变。",
            )
            unchanged = call("GET", prefix + f"/tasks/{task['id']}")
            assert unchanged["version"] == task["version"] and unchanged["progress"] == 20
            drafts = [s for s in suggestion["suggestions"] if s["suggestion_type"] == "update_task"]
            assert len(drafts) == 1, "需恰好一条更新草稿"
            draft = drafts[0]
            assert draft["proposed_payload"]["task_id"] == task["id"]
            submitted = call(
                "POST",
                prefix + f"/suggestions/{draft['id']}/submit",
                json={"expected_version": draft["version"]},
            )
            approved = call(
                "POST",
                prefix + f"/suggestions/{draft['id']}/review",
                json={"action": "confirm", "expected_version": submitted["version"]},
            )
            updated = call("GET", prefix + f"/tasks/{task['id']}")
            assert updated["progress"] == 70 and updated["version"] == task["version"] + 1
            history = call("GET", prefix + f"/tasks/{task['id']}/history")
            evidence["steps"].append(
                {"kind": "human_review", "approved": approved, "task": updated, "history": history}
            )
            extraction = call(
                "POST", prefix + f"/documents/{did}/extract", json={"kind": "extract"}
            )
            evidence["steps"].append({"kind": "document_extract_requested", "result": extraction})
            # 文档抽取跳过提交者预审，直接进入正式审核。等待实际 run 终态。
            run_id = extraction["workflow_run_id"]
            if run_id:
                for _ in range(120):
                    run = call("GET", prefix + f"/workflow-runs/{run_id}")
                    if run["status"] in {"succeeded", "failed"}:
                        break
                    time.sleep(1)
                evidence["steps"].append({"kind": "document_extract_result", "result": run})
                assert run["status"] == "succeeded"
            evidence["status"] = "succeeded"
        except Exception as error:
            evidence.update(status="failed", error=str(error))
            raise
        finally:
            (output / "live_workflow.json").write_text(
                json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8"
            )
    print("live document, assistant, review and task history verified")


if __name__ == "__main__":
    main()
