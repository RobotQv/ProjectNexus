"""Alpha 的边界验证：离线执行，真实网络评估另行显式运行。"""

import json
from concurrent.futures import ThreadPoolExecutor
from threading import Event

import httpx
import pytest
from docx import Document

from app.core.config import Settings
from app.core.security import issue_token
from app.llm import GLMProvider
from app.services.tools import BoundTools
from shared.contracts import AssistantResult, Block, DocumentRef, EntityRecord, FileRef, Scope
from shared.entity_matching import normalize, resolve_exact
from shared.errors import AppError
from shared.progress import emit
from team_modules.main_rag.adapter import MainRAGAdapter
from team_modules.main_rag.chunking import chunk_document
from team_modules.main_rag.parsing.docx import parse_docx
from team_modules.main_rag.retrieval import rrf, tokenize
from tests.conftest import new_task


def block(i, text, **extra):
    return Block(id=i, block_no=i, document_id=1, document_version=1, text=text, **extra)


def test_docx_body_table_body_order(tmp_path):
    doc = Document()
    doc.add_paragraph("表格前")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "任务", "进度"
    table.cell(1, 0).text, table.cell(1, 1).text = "登录接口", "70%"
    doc.add_paragraph("表格后")
    path = tmp_path / "table.docx"
    doc.save(path)
    result = parse_docx(
        FileRef(
            project_id=1,
            document_id=1,
            version=1,
            path=path,
            filename=path.name,
            content_hash="test",
        )
    )
    assert [b.text for b in result.blocks] == ["表格前", "任务 | 进度", "登录接口 | 70%", "表格后"]
    assert result.blocks[2].locator == "table:1/row:2"
    assert all(b.page is None for b in result.blocks)


def test_chunk_bounds_and_reconstructable_quotes():
    blocks = [block(1, "甲" * 2401), block(2, "尾段"), block(3, "新标题", heading="新标题")]
    chunks, trace = chunk_document(blocks)
    by_id = {b.id: b.text for b in blocks}
    assert max(len(c.text) for c in chunks) <= 1200
    assert trace["length_unit"] == "characters"
    for chunk in chunks:
        assert chunk.text == "\n".join(
            by_id[s["block_id"]][s["start"] : s["end"]] for s in chunk.spans
        )
        assert chunk.text in "\n".join(by_id[s["block_id"]] for s in chunk.spans)
    assert len(chunks[-1].spans) == 1
    assert len(chunk_document([block(1, "短"), block(2, "段")])[0]) == 1


def test_entity_exact_does_not_erase_punctuation_or_hide_collisions():
    def entity(i, title, aliases=[]):
        return EntityRecord(
            project_id=1,
            entity_type="task",
            entity_id=i,
            source_version=1,
            search_text=title,
            title=title,
            aliases=aliases,
            is_active=True,
        )

    catalog = [entity(1, "C++ 1.2", ["接口"]), entity(2, "C# 12", ["接口"])]
    assert normalize("C++ 1.2") != normalize("C# 12")
    assert resolve_exact(catalog, "请把C++ 1.2的进度改为70%").candidates[0].entity_id == 1
    assert resolve_exact(catalog, "接口", limit=1).outcome == "ambiguous"
    assert resolve_exact(catalog, "TASK-999 进度").outcome == "not_found"
    assert resolve_exact(catalog, "C+") is None


def test_authoritative_exact_resolution_avoids_embedding(env):
    task = new_task(env, title="登录接口", aliases=["认证"])

    def unavailable(*args):
        raise AssertionError("精确目标不应调用 embedding")

    env.app.state.modules.entities.resolve = unavailable
    tools = BoundTools(env.app.state.sessions, env.pid, 1, env.app.state.modules)
    assert tools.resolve("把认证进度改为70%").candidates[0].entity_id == task["id"]
    assert tools.resolve("TASK-999 进度").outcome == "not_found"


def test_rrf_and_technical_tokens():
    assert rrf([["a", "b"], ["b", "a"]])["a"] == 1 / 61 + 1 / 62
    assert "c++" in tokenize("C++ TASK-41 v1.2 登录接口")


class FakeEmbedding:
    def embed(self, texts):
        return [[float("登录" in t) + 0.1, float("支付" in t) + 0.1] for t in texts]


def test_persisted_bm25_delete_isolation_and_index_namespaces(tmp_path):
    a = MainRAGAdapter(store_dir=tmp_path, retrieval_mode="hybrid_rrf", chunk_strategy="bounded")
    a._embedding = FakeEmbedding()
    a.ingest(
        DocumentRef(project_id=1, document_id=1, version=1, blocks=[block(1, "登录接口重试三次")])
    )
    b = MainRAGAdapter(store_dir=tmp_path, retrieval_mode="bm25", chunk_strategy="bounded")
    assert b.retrieve(Scope(project_id=1), "登录", 5)[0].block_ids == [1]
    assert not b.retrieve(Scope(project_id=2), "登录", 5)
    assert not MainRAGAdapter(store_dir=tmp_path, retrieval_mode="bm25").retrieve(
        Scope(project_id=1), "登录", 5
    )
    a.delete(Scope(project_id=1), 1, 1)
    assert not b.retrieve(Scope(project_id=1), "登录", 5)


def test_semantic_trace_is_measured_and_bounded():
    chunks, trace = chunk_document(
        [block(1, "登录"), block(2, "支付")], strategy="semantic", embedding=FakeEmbedding()
    )
    assert len(chunks) == 2
    assert trace["segments"][1]["similarity_previous"] < 0.55


@pytest.mark.parametrize("model", ["glm-4.5-air", "glm-4.7", "glm-5.3-flash", "glm-5.3"])
def test_model_capability_payload(model):
    captured = []

    def reply(request):
        captured.append(json.loads(request.content))
        return httpx.Response(200, json={"choices": [{"message": {"content": "query"}}]})

    settings = Settings(_env_file=None, jwt_secret="s" * 40, llm_api_key="fake", llm_model=model)
    provider = GLMProvider(settings, transport=httpx.MockTransport(reply))
    try:
        provider.complete([{"role": "user", "content": "查询"}], max_tokens=16)
    finally:
        provider.close()
    assert captured[0]["thinking"]["type"] == (
        "enabled" if model.startswith("glm-5.3") else "disabled"
    )
    if model.startswith("glm-5.3"):
        assert captured[0]["max_tokens"] == 4096
        assert captured[0]["reasoning_effort"] == "low"


def test_progress_visible_during_request_and_scoped(env):
    entered, finish = Event(), Event()

    def respond(request, tools):
        emit("interpreting")
        entered.set()
        assert finish.wait(10)
        emit("generating")
        return AssistantResult(answer="已回答", model_id="fake", prompt_version="test")

    env.app.state.modules.workflow.respond = respond
    key = "alpha-progress-test"
    with ThreadPoolExecutor() as pool:
        future = pool.submit(
            env.client.post,
            env.prefix + "/assistant/messages",
            json={"entry": "project_assistant", "text": "测试", "request_key": key},
        )
        try:
            assert entered.wait(10)
            path = env.prefix + f"/assistant/progress?request_key={key}"
            state = env.client.get(path).json()
            assert state["status"] == "running"
            assert [e["stage"] for e in state["events"]] == ["accepted", "interpreting"]
            env.client.post(env.prefix + "/members", json={"login_name": "member"})
            other = {"Authorization": "Bearer " + issue_token(2, env.settings)}
            assert env.client.get(path, headers=other).status_code == 404
        finally:
            finish.set()
        assert future.result().status_code == 200
    final = env.client.get(path).json()
    assert final["terminal"] and final["events"][-1]["stage"] == "completed"
    assert [e["stage"] for e in final["events"]] == [
        "accepted",
        "interpreting",
        "generating",
        "validating",
        "persisting",
        "completed",
    ]
    assert [e["seq"] for e in final["events"]] == list(range(1, len(final["events"]) + 1))


def test_failed_run_has_terminal_event_and_no_automatic_retry(env):
    calls = []

    def fail(*args):
        calls.append(1)
        emit("interpreting")
        raise AppError("llm_timeout", "超时", 504)

    env.app.state.modules.workflow.respond = fail
    payload = {"entry": "project_assistant", "text": "问题", "request_key": "fail-alpha"}
    assert env.client.post(env.prefix + "/assistant/messages", json=payload).status_code == 504
    assert env.client.post(env.prefix + "/assistant/messages", json=payload).status_code == 409
    result = env.client.get(
        env.prefix + "/assistant/progress", params={"request_key": "fail-alpha"}
    ).json()
    assert result["status"] == "failed" and result["events"][-1]["stage"] == "failed"
    assert [e["stage"] for e in result["events"]] == ["accepted", "interpreting", "failed"]
    assert len(calls) == 1


def test_embedding_batches_and_restores_index_order():
    from team_modules.main_rag.indexing.embedding import ZhipuEmbedding

    counts = []

    def reply(request):
        inputs = json.loads(request.content)["input"]
        counts.append(len(inputs))
        return httpx.Response(
            200,
            json={
                "data": [
                    {"index": i, "embedding": [float(text)] * 2048}
                    for i, text in reversed(list(enumerate(inputs)))
                ]
            },
        )

    provider = ZhipuEmbedding(api_key="fake")
    provider.client.close()
    provider.client = httpx.Client(transport=httpx.MockTransport(reply))
    try:
        result = provider.embed([str(i) for i in range(35)])
    finally:
        provider.close()
    assert counts == [16, 16, 3]
    assert [v[0] for v in result] == list(range(35))


def test_same_origin_spa_does_not_mask_api_or_expose_parent_files(env, tmp_path):
    from fastapi.testclient import TestClient

    from app.main import create_app

    dist = tmp_path / "web"
    dist.mkdir()
    (dist / "index.html").write_text("<html>Alpha app</html>")
    (tmp_path / "secret.txt").write_text("must-not-be-public")
    app = create_app(env.settings.model_copy(update={"web_dist_dir": dist}))
    with TestClient(app) as client:
        assert client.get("/").text == "<html>Alpha app</html>"
        assert client.get("/some-spa-route").text == "<html>Alpha app</html>"
        assert client.get("/api/v1/does-not-exist").status_code == 404
        assert "must-not-be-public" not in client.get("/%2e%2e/secret.txt").text
        assert client.get("/api/v1/projects").status_code == 401
