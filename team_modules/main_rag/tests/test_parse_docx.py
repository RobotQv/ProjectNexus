"""DOCX 解析测试。"""

from pathlib import Path

import pytest
from docx import Document

from shared.contracts import FileRef
from shared.errors import AppError
from team_modules.main_rag.adapter import MainRAGAdapter


def make_docx(tmp_path: Path, paragraphs: list[tuple[str, str]]) -> FileRef:
    """paragraphs: [(text, style), ...]，style 为空表示普通段落。"""
    path = tmp_path / "test.docx"
    doc = Document()
    for text, style in paragraphs:
        if style:
            doc.add_paragraph(text, style=style)
        else:
            doc.add_paragraph(text)
    doc.save(str(path))
    return FileRef(
        project_id=1,
        document_id=1,
        version=1,
        path=path,
        filename="test.docx",
        content_hash="abc123",
    )


def test_parse_docx_basic(tmp_path: Path):
    file = make_docx(
        tmp_path,
        [
            ("标题一", "Heading 1"),
            ("第一段。", ""),
            ("标题二", "Heading 2"),
            ("第二段。", ""),
        ],
    )
    doc = MainRAGAdapter().parse(file)
    assert len(doc.blocks) == 4
    assert doc.blocks[0].text == "标题一"
    assert doc.blocks[0].heading == "标题一"
    assert doc.blocks[1].text == "第一段。"
    assert doc.blocks[1].heading == "标题一"
    assert doc.blocks[2].text == "标题二"
    assert doc.blocks[3].text == "第二段。"
    assert doc.blocks[3].heading == "标题二"


def test_parse_docx_empty(tmp_path: Path):
    file = make_docx(tmp_path, [("", "")])
    with pytest.raises(AppError) as exc:
        MainRAGAdapter().parse(file)
    assert exc.value.code == "empty_document"


def test_parse_docx_missing(tmp_path: Path):
    file = FileRef(
        project_id=1,
        document_id=1,
        version=1,
        path=tmp_path / "not_exist.docx",
        filename="not_exist.docx",
        content_hash="abc123",
    )
    with pytest.raises(AppError) as exc:
        MainRAGAdapter().parse(file)
    assert exc.value.code == "file_not_found"


def test_parse_docx_too_large(tmp_path: Path, monkeypatch):
    file = make_docx(tmp_path, [("内容", "")])
    monkeypatch.setattr(
        "team_modules.main_rag.parsing.docx.MAX_DOCX_BYTES",
        1,
    )
    with pytest.raises(AppError) as exc:
        MainRAGAdapter().parse(file)
    assert exc.value.code == "document_too_large"


def test_parse_docx_no_page(tmp_path: Path):
    file = make_docx(tmp_path, [("标题", "Heading 1"), ("正文。", "")])
    doc = MainRAGAdapter().parse(file)
    assert all(block.page is None for block in doc.blocks)
